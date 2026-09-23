"""Generise upitnik za klijenta: docs/Dr-Cooker-upitnik.docx

Kratak dokument (do 10 strana), bez uvoda, sa checkbox-ovima koji se STVARNO
klikcu u Wordu.

Zasto Word content control a ne znak ☐:
  Znak ☐ je obican tekst i ne moze da se klikne. Legacy form field
  (FORMCHECKBOX) moze, ali trazi da dokument bude zakljucan za forme, sto bi
  onemogucilo kucanje slobodnog teksta u ostatku dokumenta. Content control
  (w14:checkbox) se klikce bez zakljucavanja, a ostatak dokumenta ostaje
  normalno izmenjiv. Radi u Wordu 2010 i novijim.

Pokretanje:
    pip install python-docx
    python tools/build-docx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "Dr-Cooker-upitnik.docx"

INK = RGBColor(0x17, 0x14, 0x0F)
SOFT = RGBColor(0x4A, 0x42, 0x38)
ACCENT = RGBColor(0xA3, 0x2A, 0x2E)
PROOF = RGBColor(0x12, 0x4A, 0x3F)
RULE = "8C7E69"
TINT = "E6EDE8"

doc = Document()

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(9.5)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(3)
st.paragraph_format.line_spacing = 1.05

for s in doc.sections:
    s.top_margin = Cm(1.5)
    s.bottom_margin = Cm(1.3)
    s.left_margin = Cm(1.7)
    s.right_margin = Cm(1.7)

_uid = [100000]


def _next_id() -> str:
    _uid[0] += 1
    return str(_uid[0])


def _shade(el, color: str) -> None:
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), color)
    el.append(sh)


def _border(par, edge: str, color: str = RULE, size: int = 6) -> None:
    pPr = par._p.get_or_add_pPr()
    b = pPr.find(qn("w:pBdr"))
    if b is None:
        b = OxmlElement("w:pBdr")
        pPr.append(b)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(size))
    e.set(qn("w:space"), "2")
    e.set(qn("w:color"), color)
    b.append(e)


def _fixed(t) -> None:
    """Bez ovoga Word ignorise zadate sirine kolona."""
    t.autofit = False
    pr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    pr.append(lay)
    w = OxmlElement("w:tblW")
    w.set(qn("w:type"), "pct")
    w.set(qn("w:w"), "5000")
    pr.append(w)


def _row_h(row, cm: float) -> None:
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


# ---------------------------------------------------------------------------
#  KLIKABILAN CHECKBOX — Word content control
# ---------------------------------------------------------------------------


def add_checkbox(par) -> None:
    """Ubacuje kvadratic koji se u Wordu klikom prebacuje u kvacicu."""
    sdt = OxmlElement("w:sdt")

    pr = OxmlElement("w:sdtPr")
    sid = OxmlElement("w:id")
    sid.set(qn("w:val"), _next_id())
    pr.append(sid)

    cb = OxmlElement("w14:checkbox")
    checked = OxmlElement("w14:checked")
    checked.set(qn("w14:val"), "0")
    on = OxmlElement("w14:checkedState")
    on.set(qn("w14:val"), "2612")  # ☒
    on.set(qn("w14:font"), "MS Gothic")
    off = OxmlElement("w14:uncheckedState")
    off.set(qn("w14:val"), "2610")  # ☐
    off.set(qn("w14:font"), "MS Gothic")
    cb.append(checked)
    cb.append(on)
    cb.append(off)
    pr.append(cb)
    sdt.append(pr)

    content = OxmlElement("w:sdtContent")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        f.set(qn(a), "MS Gothic")
    rPr.append(f)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = chr(9744)
    r.append(t)
    content.append(r)
    sdt.append(content)

    par._p.append(sdt)


# ---------------------------------------------------------------------------
#  Gradivni elementi
# ---------------------------------------------------------------------------


def section(letter: str, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(letter + "  ")
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    r = p.add_run(title.upper())
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = PROOF
    _border(p, "bottom", "124A3F", 8)


def note(text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(8)
    r.font.italic = True
    r.font.color.rgb = SOFT


def table(headers: list[str], rows: int, widths: list[float], h: float = 0.72) -> None:
    t = doc.add_table(rows=rows + 1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, head in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(head)
        r.font.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c._tc.get_or_add_tcPr(), "124A3F")
        c.paragraphs[0].paragraph_format.space_after = Pt(1)
    _fixed(t)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    _row_h(t.rows[0], 0.5)
    for row in t.rows[1:]:
        _row_h(row, h)
        for c in row.cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def checks(items: list[str], cols: int = 2) -> None:
    """Klikabilni checkbox-ovi u vise kolona, da ne trose visinu strane."""
    rows = (len(items) + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols)
    _fixed(t)
    width = 16.6 / cols
    for i, item in enumerate(items):
        c = t.rows[i % rows].cells[i // rows]
        c.width = Cm(width)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        add_checkbox(p)
        r = p.add_run("  " + item)
        r.font.size = Pt(9)
    for i in range(len(items), rows * cols):
        t.rows[i % rows].cells[i // rows].width = Cm(width)
    # bez vidljivih linija — ovo je lista, ne tabela
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def lines(labels: list[str], per_row: int = 2) -> None:
    """Kratka polja: labela pa linija za upis, vise u redu."""
    t = doc.add_table(rows=(len(labels) + per_row - 1) // per_row, cols=per_row)
    _fixed(t)
    width = 16.6 / per_row
    for i, lab in enumerate(labels):
        c = t.rows[i // per_row].cells[i % per_row]
        c.width = Cm(width)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(lab)
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = SOFT
        a = c.add_paragraph()
        a.paragraph_format.space_before = Pt(7)
        a.paragraph_format.space_after = Pt(3)
        _border(a, "bottom", RULE, 4)
    for i in range(len(labels), len(t.rows) * per_row):
        t.rows[i // per_row].cells[i % per_row].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def box(rows: int = 3, label: str = "") -> None:
    if label:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label)
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = SOFT
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    _fixed(t)
    _row_h(t.rows[0], 0.52 * rows)
    t.rows[0].cells[0].paragraphs[0].paragraph_format.space_after = Pt(1)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


# ===========================================================================
#  ZAGLAVLJE
# ===========================================================================

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Dr Cooker")
r.font.size = Pt(19)
r.font.bold = True
r.font.color.rgb = PROOF
r = p.add_run("   Upitnik za sadržaj sajta")
r.font.size = Pt(13)
r.font.color.rgb = INK

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
_border(p, "bottom", "A32A2E", 14)
r = p.add_run(
    "Popunite samo ono što se odnosi na Vas. Prazno polje znači da ta informacija "
    "neće stajati na sajtu — ništa ne dopisujemo."
)
r.font.size = Pt(8.5)
r.font.color.rgb = SOFT

t = doc.add_table(rows=1, cols=3)
_fixed(t)
for i, lab in enumerate(["Popunio/la:", "Funkcija:", "Datum:"]):
    c = t.rows[0].cells[i]
    c.width = Cm(5.53)
    pp = c.paragraphs[0]
    pp.paragraph_format.space_after = Pt(0)
    rr = pp.add_run(lab)
    rr.font.size = Pt(8)
    rr.font.bold = True
    rr.font.color.rgb = SOFT
    a = c.add_paragraph()
    a.paragraph_format.space_before = Pt(7)
    _border(a, "bottom", RULE, 4)
doc.add_paragraph().paragraph_format.space_after = Pt(1)

# ===========================================================================

section("A", "Kompanija")
lines(["Pun pravni naziv", "PIB i matični broj", "Adresa kuhinje", "Adresa kancelarije"], 2)
lines(["Godina početka rada", "Radno vreme (telefonom)", "Domen koji želite", "Kontakt osoba za sajt"], 2)
note(
    "Na sajtu sada piše da je kuhinja „na Voždovcu“, a kontakt adresa je Vojvode "
    "Prijezde 17. Da li je to ista adresa?"
)

section("B", "O firmi")
lines(["Koliko godina iskustva ima tim", "Broj zaposlenih", "Broj kuvara", "Broj vozila"], 4)
lines(["Maks. obroka dnevno", "Trenutno obroka dnevno", "Kvadratura kuhinje", "Broj ustanova koje snabdevate"], 4)
box(3, "Po čemu se razlikujete od druge ketering firme u Beogradu — konkretno, ne „kvalitet i posvećenost“")

section("C", "Sertifikati")
note("Sajt sada kaže „poštuju se principi HACCP standarda“. To NIJE isto što i sertifikat. Ako ga imate, to je najjači dokaz koji možete da pokažete.")
checks(
    [
        "Imamo HACCP sertifikat (šaljemo kopiju)",
        "Primenjujemo principe, nemamo sertifikat",
        "Imamo ISO sertifikat",
        "Imamo rešenje sanitarne inspekcije",
    ],
    2,
)
lines(["Broj i izdavalac sertifikata", "Važi do"], 2)

section("D", "Ciljne grupe")
table(
    ["Grupa", "Radite?", "Broj klijenata", "Tipična veličina", "Šta im nudite"],
    5,
    [3.4, 1.6, 2.4, 2.6, 6.6],
)
note("Redom: privatni vrtići i jaslice · boravci za decu · privatne škole.")

section("E", "Usluge i cene")
table(
    ["Usluga", "Kome je namenjena", "Cena", "Način obračuna", "Minimum", "Rok narudžbine"],
    5,
    [3.4, 3.4, 2.2, 2.6, 2.4, 2.6],
)

section("F", "Paketi")
checks(["Imamo definisane pakete", "Nemamo — svaka ponuda je posebna", "Želimo da ih napravimo"], 3)
table(
    ["Naziv paketa", "Za koga", "Broj osoba", "Šta uključuje", "Cena po osobi"],
    5,
    [3.0, 2.6, 2.0, 6.4, 2.6],
)

section("G", "Meni")
lines(["Ko sastavlja meni", "Na koliko nedelja se ponavlja", "Menja li se po sezoni", "Koliko unapred ga klijent dobija"], 2)
table(["Kategorija", "Primeri jela", "Gramaža", "Cena", "Alergeni"], 5, [3.0, 6.2, 2.0, 2.0, 3.4])
note("Ako imate meni u Excelu ili Wordu — pošaljite ga umesto prekucavanja. Kategorije: doručak, užine, supe, glavna jela, salate, deserti.")

section("H", "Posebni režimi ishrane")
checks(
    ["Bez glutena", "Bez laktoze", "Vegetarijansko", "Vegansko", "Bez orašastih plodova", "Dijabetička ishrana", "Posna hrana", "Drugo"],
    4,
)
lines(["Doplaćuje li se i koliko", "Koliko unapred se prijavljuje"], 2)

section("I", "Vrtići, jaslice i škole")
table(["Obrok", "Vreme isporuke", "Cena po detetu", "Napomena"], 4, [3.4, 3.4, 3.4, 6.4])
lines(["Min. broj dece za ugovor", "Maks. broj dece", "Kako se obračunava odsustvo deteta", "Ko preuzima hranu u ustanovi"], 2)
box(2, "Kako izgleda prijava alergije i zamena namirnice")

section("J", "Boravci za decu")
note("Boravci su jedna od tri publike koje ste naveli, ali o njima na starom sajtu nema nijednog podatka. Forma sada nudi „Kuvani obroci — boravak za decu“ kao zasebnu opciju.")
table(["Obrok", "Vreme isporuke", "Cena po detetu", "Napomena"], 4, [3.4, 3.4, 3.4, 6.4])
lines(["Po čemu se posao za boravak razlikuje od vrtića", "Min. broj dece za ugovor"], 2)

section("K", "Dostava")
note("Sajt trenutno ne navodi područje dostave — te informacije nigde nema. To je jedno od najčešćih pitanja pre slanja upita.")
lines(["Opštine / gradovi koje pokrivate", "Izlazite li van Beograda"], 2)
table(["Zona", "Cena dostave", "Min. narudžbina", "Termin isporuke"], 4, [4.6, 3.6, 4.0, 4.4])
checks(["Dostavljamo vikendom", "Dostavljamo praznicima", "Dostava besplatna preko iznosa:"], 3)

section("L", "Plaćanje")
checks(["Faktura sa odloženim plaćanjem", "Avansno", "Po isporuci", "Gotovina", "Kartica"], 3)
lines(["Rok plaćanja za pravna lica", "Visina avansa", "Do kada se otkazuje bez naplate", "Popust za duži ugovor"], 2)

section("M", "Reference")
note("Sajt sada pominje „najveću privatnu predškolsku ustanovu u Srbiji“ bez imena, jer stari sajt ne daje ime. Imenovana referenca vredi višestruko više.")
table(["Klijent", "Tip", "Od kada", "Obim (broj korisnika)", "Sme ime javno?"], 5, [4.4, 2.6, 2.0, 3.6, 4.0])

section("N", "Preporuke klijenata")
note("Sajt nema nijednu, jer je stari sajt nema. Preporuka direktora vrtića je najjači argument koji možete imati. Potrebna je stvarna izjava i dozvola za objavu.")
table(["Ime i funkcija", "Ustanova", "Tekst preporuke", "Dozvola?"], 3, [3.4, 3.0, 7.6, 2.6])

section("O", "Fotografije — najvažnije")
note(
    "Na starom sajtu ne postoji nijedna prava fotografija Vaše hrane, kuhinje, tima "
    "ni dostave — sve su kupljene stock slike. Na sajtu je 19 mesta već pripremljeno; "
    "čim pošaljete slike, ubacujemo ih bez izmene rasporeda. Telefon i dnevno svetlo su dovoljni."
)
checks(
    [
        "Kuhinja tokom pripreme (3–5)",
        "Serviran obrok odozgo (5–8)",
        "Vozilo i termo posude (2–3)",
        "Tim u kuhinji (1–2)",
        "Prijem namirnica (2)",
        "Pakovanje obroka (2–3)",
        "Nutricionista na poslu (2)",
        "Obrok u vrtiću ili školi (3–5)",
        "Logo u vektoru (AI, EPS, SVG, PDF)",
        "Nemamo ništa od navedenog",
    ],
    2,
)
lines(["Kada možete da snimite", "Ko šalje fotografije"], 2)

section("P", "Brend")
note("Vaš postojeći logo je vraćen u zaglavlje i podnožje, po dogovoru — ne menja se. Jedino što treba je bolji fajl: dostavljena verzija je 260×64 piksela, što je premalo za ekrane sa dvostrukom gustinom, pa je logo na njima blago mek.")
checks(
    ["Šaljemo vektor (SVG, AI, EPS ili PDF)", "Šaljemo PNG od bar 1000 px širine", "Nemamo ništa veće od postojećeg"],
    1,
)
lines(["Slogan koji koristite", "Boje brenda (ako su definisane)"], 2)

section("Q", "Kontakt i forma")
table(["Podatak", "Sada na sajtu", "Tačno / ispraviti na"], 1, [4.4, 5.6, 6.6])
_t = doc.tables[-1]
for podatak, sada in [
    ("Telefon — direktor", "+381 64 1101521"),
    ("Telefon — komercijala", "+381 60 0440458"),
    ("Email", "office@drcooker.rs"),
    ("Adresa", "Vojvode Prijezde 17, Beograd"),
    ("Društvene mreže", "nisu navedene"),
]:
    row = _t.add_row()
    _row_h(row, 0.62)
    for i, val in enumerate([podatak, sada, ""]):
        row.cells[i].text = ""
        rr = row.cells[i].paragraphs[0].add_run(val)
        rr.font.size = Pt(8.5)
        row.cells[i].paragraphs[0].paragraph_format.space_after = Pt(1)
        if i == 0:
            rr.font.bold = True
        if i == 1:
            rr.font.color.rgb = SOFT
doc.add_paragraph().paragraph_format.space_after = Pt(1)
lines(["Ko prima upite sa sajta (email)", "U kom roku odgovarate na upit"], 2)
note("Rok koji upišete ide na sajt i u automatsku potvrdu klijentu. Upišite onaj koji sigurno možete da ispunite; ako niste sigurni, ostavite prazno.")

section("R", "Sajt")
checks(
    ["Galerija fotografija", "Cenovnik", "Blog / saveti o ishrani", "Jelovnik za preuzimanje (PDF)", "Stranica za događaje", "Onlajn poručivanje"],
    3,
)
box(2, "Šta Vam najviše nedostaje na sajtu i šta želite da postignete u sledećih godinu dana")

section("S", "Česta pitanja (FAQ)")
note(
    "Sajt već ima blokove čestih pitanja na stranicama za ustanove i "
    "snabdevanje. Odgovori su sada izvedeni iz starog sajta. Ovde upišite pitanja "
    "koja vam klijenti STVARNO najčešće postavljaju — ona vrede više od bilo kog "
    "teksta koji bismo mi smislili."
)
table(["Pitanje koje vam najčešće postavljaju", "Vaš odgovor"], 6, [6.6, 10.0], 0.78)
checks(
    [
        "Pitaju za cenu pre svega",
        "Pitaju za alergije i posebne režime",
        "Pitaju za probni period",
        "Pitaju za minimalni broj korisnika",
        "Pitaju ko kontroliše kvalitet",
        "Pitaju za otkazivanje ugovora",
    ],
    3,
)

section("T", "Sinko online prodavnica (sinko.rs)")
note(
    "Sajt sada upućuje klijente na sinko.rs. Sve što na njemu piše o Sinku uzeto je "
    "sa samog sinko.rs. Ono što sa tog sajta NE može da se pročita je odnos dve firme "
    "\u2014 Sinko plus d.o.o. ima drugi PIB i drugu adresu. Dok ovo ne potvrdite, sajt "
    "koristi neutralnu reč \u201epartner\u201c."
)
checks(
    [
        "Ista vlasnička struktura",
        "Poslovni partner",
        "Sinko je naš dobavljač",
        "Mi smo Sinkov dobavljač",
        "Nema formalne veze",
    ],
    3,
)
box(2, "Tačna rečenica kojom sajt sme da opiše odnos Dr Cooker-a i Sinko plus d.o.o.")
table(
    ["Pitanje", "Odgovor"],
    6,
    [7.2, 9.4],
    0.80,
)
note(
    "Redom: 1) Da li je asortiman na sinko.rs isti kao onih 7000 artikala koje već "
    "isporučujete? \u00b7 2) Sme li sajt da tvrdi da je online cena niža od cene u radnji? "
    "\u00b7 3) Koji broj klijent zove za porudžbinu sa sinko.rs? \u00b7 4) Da li pragovi "
    "(4.999 / 7.999 RSD) i zona do 10 km i dalje važe? \u00b7 5) Kako plaćaju pravna lica "
    "\u2014 samo virmanski ili i drugim načinima? \u00b7 6) Dobijaju li vaši klijenti "
    "keteringa nešto posebno (nalog, rabat, zbirna faktura)?"
)

# --- zavrsetak --------------------------------------------------------------

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
_border(p, "top", "124A3F", 8)
r = p.add_run(
    "Ne morate popuniti sve odjednom — pošaljite delove kako budu gotovi. "
    "Fotografije i logo šaljite kao posebne fajlove."
)
r.font.size = Pt(8.5)
r.font.color.rgb = SOFT

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)

# .docx je ZIP — sadrzaj se mora raspakovati da bi se prebrojao.
import zipfile

with zipfile.ZipFile(OUT) as z:
    xml = z.read("word/document.xml").decode("utf-8")
cb_count = xml.count("w14:checkbox>") // 2  # otvarajuci + zatvarajuci tag
print(f"Napisano: {OUT.relative_to(ROOT)}")
print(f"Velicina: {OUT.stat().st_size / 1024:.1f} KB")
print(f"Tabela:   {len(doc.tables)}")
print(f"Klikabilnih checkbox-ova: {cb_count}")
print(f"Znakova ☐ kao obican tekst: {sum(p.text.count(chr(9744)) for p in doc.paragraphs)}")
